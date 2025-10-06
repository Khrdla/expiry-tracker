#!/usr/bin/env python3
"""
Simple Enhanced Inventory Scanning Excel Export Test
Focus on core multi-zone worksheet functionality
"""

import asyncio
import aiohttp
import json
import openpyxl
from io import BytesIO

# Configuration
BACKEND_URL = "https://geant-scanner.preview.emergentagent.com/api"
ADMIN_USERNAME = "imadqejji"
ADMIN_PASSWORD = "066380531I"

async def test_inventory_scanning_export():
    """Test the Enhanced Inventory Scanning Excel Export with multiple worksheets"""
    
    print("🚀 TESTING ENHANCED INVENTORY SCANNING EXCEL EXPORT")
    print("=" * 70)
    
    session = aiohttp.ClientSession()
    
    try:
        # Step 1: Authenticate
        print("1. Authenticating...")
        login_data = {"username": ADMIN_USERNAME, "password": ADMIN_PASSWORD}
        async with session.post(f"{BACKEND_URL}/auth/login", json=login_data) as response:
            if response.status != 200:
                print(f"❌ Authentication failed: {response.status}")
                return False
            data = await response.json()
            auth_token = data.get("access_token")
            headers = {"Authorization": f"Bearer {auth_token}"}
            print("✅ Authentication successful")
        
        # Step 2: Clear existing scans
        print("\n2. Clearing existing scans...")
        async with session.delete(f"{BACKEND_URL}/inventory-scans/clear", headers=headers) as response:
            if response.status == 200:
                data = await response.json()
                print(f"✅ {data.get('message', 'Cleared scans')}")
            else:
                print(f"⚠️ Clear failed: {response.status}")
        
        # Step 3: Create test scans in multiple zones
        print("\n3. Creating inventory scans in multiple zones...")
        test_scans = [
            ("SA", 1, "3222471081716", 50),  # Apple Juice Box 1L in SA01
            ("SA", 2, "3222471052747", 30),  # Lemonade 150Cl in SA02
            ("WH", 1, "3222471075722", 25),  # Mountain Water 6X50Cl in WH01
            ("WH", 2, "3222471081273", 40),  # Orange Peach Apricot Nectar in WH02
        ]
        
        created_scans = 0
        for zone_type, zone_number, barcode, quantity in test_scans:
            scan_data = {
                "zone_type": zone_type,
                "zone_number": zone_number,
                "barcode": barcode,
                "quantity_scanned": quantity
            }
            
            async with session.post(f"{BACKEND_URL}/inventory-scans", json=scan_data, headers=headers) as response:
                if response.status == 200:
                    created_scans += 1
                    print(f"✅ Created scan: {zone_type}{zone_number:02d} - {barcode} - {quantity} qty")
                else:
                    error_text = await response.text()
                    print(f"❌ Failed to create scan {zone_type}{zone_number:02d}: {response.status} - {error_text}")
        
        print(f"✅ Created {created_scans}/{len(test_scans)} inventory scans")
        
        if created_scans == 0:
            print("❌ No scans created - cannot test export")
            return False
        
        # Step 4: Export Excel and analyze
        print("\n4. Exporting Excel file...")
        async with session.get(f"{BACKEND_URL}/inventory-scans/export", headers=headers) as response:
            if response.status != 200:
                error_text = await response.text()
                print(f"❌ Export failed: {response.status} - {error_text}")
                return False
            
            excel_content = await response.read()
            print(f"✅ Excel file downloaded: {len(excel_content)} bytes")
            
            # Verify Content-Type
            content_type = response.headers.get('content-type', '')
            if 'spreadsheet' in content_type:
                print("✅ Correct Content-Type for Excel file")
            else:
                print(f"⚠️ Unexpected Content-Type: {content_type}")
            
            # Verify filename
            content_disposition = response.headers.get('content-disposition', '')
            if 'Inventory_Scan_Report_' in content_disposition:
                print("✅ Proper filename format with timestamp")
            else:
                print(f"⚠️ Unexpected filename: {content_disposition}")
        
        # Step 5: Analyze Excel structure
        print("\n5. Analyzing Excel structure...")
        workbook = openpyxl.load_workbook(BytesIO(excel_content))
        worksheet_names = workbook.sheetnames
        print(f"✅ Found {len(worksheet_names)} worksheets: {worksheet_names}")
        
        # Check for expected zones
        expected_zones = ["SA01", "SA02", "WH01", "WH02"]
        zones_found = []
        for expected_zone in expected_zones:
            if expected_zone in worksheet_names:
                zones_found.append(expected_zone)
                print(f"✅ Worksheet found for zone: {expected_zone}")
            else:
                print(f"❌ Missing worksheet for zone: {expected_zone}")
        
        # Step 6: Analyze individual worksheets
        print("\n6. Analyzing individual worksheets...")
        for zone_name in zones_found:
            worksheet = workbook[zone_name]
            print(f"\n📋 WORKSHEET: {zone_name}")
            
            # Check zone header
            zone_header = worksheet['A1'].value
            if zone_header and f"Zone Number: {zone_name}" in str(zone_header):
                print(f"✅ Zone header: {zone_header}")
            else:
                print(f"❌ Zone header incorrect: {zone_header}")
            
            # Check SKU count
            sku_count_cell = worksheet['A2'].value
            if sku_count_cell and "Total SKUs Scanned:" in str(sku_count_cell):
                print(f"✅ SKU count: {sku_count_cell}")
            else:
                print(f"❌ SKU count missing: {sku_count_cell}")
            
            # Check timestamp
            timestamp_cell = worksheet['A3'].value
            if timestamp_cell and "Generated On:" in str(timestamp_cell):
                print(f"✅ Timestamp: {timestamp_cell}")
            else:
                print(f"❌ Timestamp missing: {timestamp_cell}")
            
            # Check table headers (row 5)
            expected_headers = [
                "Item Number", "Barcode", "Description", "Supplier Code", "Supplier Name",
                "Qty Scanned in SA", "Qty Scanned in WH", "Total Inventory Scan",
                "System Stock", "Variance in Qty", "Variance in Value"
            ]
            
            headers_correct = True
            for col_num, expected_header in enumerate(expected_headers, 1):
                actual_header = worksheet.cell(row=5, column=col_num).value
                if actual_header != expected_header:
                    print(f"❌ Header mismatch col {col_num}: expected '{expected_header}', got '{actual_header}'")
                    headers_correct = False
            
            if headers_correct:
                print("✅ All table headers correct")
            
            # Check header formatting
            header_cell = worksheet.cell(row=5, column=1)
            if header_cell.font.bold:
                print("✅ Headers are bold")
            else:
                print("❌ Headers are not bold")
            
            # Check frozen panes
            if worksheet.freeze_panes:
                print(f"✅ Frozen panes: {worksheet.freeze_panes}")
            else:
                print("❌ No frozen panes")
            
            # Count data rows
            data_rows = 0
            row = 6  # First data row
            while worksheet.cell(row=row, column=1).value:
                data_rows += 1
                row += 1
            
            print(f"✅ Found {data_rows} data rows")
            
            # Look for totals row
            totals_row = 5 + data_rows + 1
            totals_label = worksheet.cell(row=totals_row, column=5).value
            if totals_label == "TOTALS:":
                print("✅ Totals row found")
                
                # Check totals values
                total_sa = worksheet.cell(row=totals_row, column=6).value
                total_wh = worksheet.cell(row=totals_row, column=7).value
                total_inventory = worksheet.cell(row=totals_row, column=8).value
                total_variance = worksheet.cell(row=totals_row, column=11).value
                
                print(f"✅ Zone totals: SA={total_sa}, WH={total_wh}, Total={total_inventory}, Variance={total_variance}")
            else:
                print(f"❌ Totals row not found. Row {totals_row}, Col 5 = '{totals_label}'")
        
        # Step 7: Summary
        print("\n" + "=" * 70)
        print("📊 TEST SUMMARY")
        print("=" * 70)
        
        success_criteria = [
            len(zones_found) == len(expected_zones),  # All zones created
            len(excel_content) > 5000,  # Reasonable file size
            'spreadsheet' in content_type,  # Correct content type
            'Inventory_Scan_Report_' in content_disposition  # Correct filename
        ]
        
        passed_criteria = sum(success_criteria)
        total_criteria = len(success_criteria)
        
        print(f"✅ PASSED: {passed_criteria}/{total_criteria} core criteria")
        print(f"📊 WORKSHEETS: {len(zones_found)}/{len(expected_zones)} zones created")
        print(f"📁 FILE SIZE: {len(excel_content)} bytes")
        
        if passed_criteria == total_criteria:
            print("🎉 ENHANCED INVENTORY SCANNING EXCEL EXPORT: FULLY FUNCTIONAL!")
            return True
        else:
            print("🚨 SOME ISSUES DETECTED - NEEDS ATTENTION")
            return False
        
    except Exception as e:
        print(f"❌ Test error: {str(e)}")
        return False
    finally:
        await session.close()

if __name__ == "__main__":
    asyncio.run(test_inventory_scanning_export())